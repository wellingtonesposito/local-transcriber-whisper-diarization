from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt

from ..core.model import Edits, Transcript
from ..core.render import ExportOptions, build_cues, fmt_ts


def _doc(transcript: Transcript) -> Document:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(transcript.media.get("filename", "Transcript"), level=1)
    return doc


def _bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_table(transcript: Transcript, edits: Edits | None, opts: ExportOptions) -> bytes:
    """Timespan | Speaker | Content — the layout NVivo's transcript importer maps columns from.

    No heading is written above the table so the file is a clean import source.
    """
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    cues = build_cues(transcript, edits, opts.unmerged() if opts.merge_turns is None else opts)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("Timespan", "Speaker", "Content")):
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
    p = opts.ts_precision
    for c in cues:
        row = table.add_row().cells
        row[0].text = f"{fmt_ts(c.start, p)}-{fmt_ts(c.end, p)}"
        row[1].text = c.speaker
        row[2].text = c.text
    return _bytes(doc)


def render_paragraphs(transcript: Transcript, edits: Edits | None, opts: ExportOptions) -> bytes:
    """One paragraph per cue: `[hh:mm:ss.s] Speaker: text`."""
    doc = _doc(transcript)
    for c in build_cues(transcript, edits, opts):
        para = doc.add_paragraph()
        if opts.timestamps:
            para.add_run(f"[{fmt_ts(c.start, opts.ts_precision)}] ")
        para.add_run(f"{c.speaker}: ").bold = True
        para.add_run(c.text)
    return _bytes(doc)
