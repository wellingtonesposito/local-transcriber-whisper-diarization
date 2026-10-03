import csv
import io
import json

import pytest
from docx import Document

from transcriber.core.model import Edits
from transcriber.core.render import ExportOptions, build_cues, fmt_ts
from transcriber.exporters import get_exporters, render


def test_fmt_ts_rounds_not_truncates():
    assert fmt_ts(1.9996, 3) == "00:00:02.000"
    assert fmt_ts(3661.5, 1) == "01:01:01.5"
    assert fmt_ts(5, 0) == "00:00:05"
    assert fmt_ts(1.5, 3, ",") == "00:00:01,500"
    assert fmt_ts(-1) == "00:00:00.000"


def test_segments_split_at_speaker_change(interview):
    assert [s.speaker for s in interview.segments[:3]] == ["SPEAKER_00", "SPEAKER_01", "SPEAKER_00"]


def test_vtt_prefix_and_voice(interview):
    out, _ = render("vtt", interview, None)
    text = out.decode()
    assert text.startswith("WEBVTT\n\n00:00:00.000 --> 00:00:02.000\nSPEAKER_00: Hello, um, thanks for joining.")
    voice, _ = render("vtt", interview, None, ExportOptions(speaker_style="voice"))
    assert "<v SPEAKER_00>Hello, um," in voice.decode()


def test_vtt_escapes_markup(interview):
    interview.segments[0].text = "a < b & c"
    assert "a &lt; b &amp; c" in render("vtt", interview, None)[0].decode()


def test_speaker_rename_and_cleanup_apply_at_export(interview):
    edits = Edits(speakers={"SPEAKER_00": "Researcher", "SPEAKER_01": "Participant"})
    txt = render("txt", interview, edits, ExportOptions(cleanup="clean"))[0].decode()
    assert "[00:00:00.0] Researcher: Hello, thanks for joining." in txt
    assert "Participant: I like pizza," in txt
    # raw transcript is untouched
    assert "um," in interview.segments[0].text


def test_segment_edit_and_delete(interview):
    edits = Edits(segments={"0": {"text": "Edited."}, "1": {"deleted": True}})
    cues = build_cues(interview, edits)
    assert cues[0].text == "Edited."
    assert all("pizza" not in c.text for c in cues)


def test_plain_txt_merges_turns_without_timestamps(interview):
    out = render("txt", interview, None, ExportOptions(timestamps=False))[0].decode()
    assert "[" not in out


def test_srt_numbering_and_comma(interview):
    out = render("srt", interview, None)[0].decode()
    assert out.startswith("1\n00:00:00,000 --> 00:00:02,000\n")


def test_csv_roundtrip(interview):
    raw, _ = render("csv", interview, None)
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0]["speaker"] == "SPEAKER_00" and float(rows[0]["end"]) == pytest.approx(2.0)


def test_json_has_words_and_rendered_text(interview):
    d = json.loads(render("json", interview, None, ExportOptions(cleanup="clean"))[0])
    assert d["segments"][0]["rendered_text"] == "Hello, thanks for joining."
    assert d["segments"][0]["words"][1]["w"] == "um,"


def test_docx_table_layout(interview):
    raw, ex = render("docx_table", interview, Edits(speakers={"SPEAKER_00": "R"}))
    doc = Document(io.BytesIO(raw))
    t = doc.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["Timespan", "Speaker", "Content"]
    assert t.rows[1].cells[0].text == "00:00:00.0-00:00:02.0"
    assert t.rows[1].cells[1].text == "R"


def test_docx_paragraphs_pdf_md_nonempty(interview):
    for key in ("docx", "pdf", "md"):
        raw, ex = render(key, interview, None)
        assert len(raw) > 100, key
    assert render("pdf", interview, None)[0].startswith(b"%PDF")


def test_all_registered_exporters_run(interview):
    for key in get_exporters():
        assert render(key, interview, None)[0]


def _two_turns():
    from transcriber.core.model import Segment, Transcript, Word

    def seg(i, s, e, spk, text):
        return Segment(i, s, e, spk, text, [Word(w, s, e) for w in text.split()])

    return Transcript(media={"filename": "m.wav"}, language="en", engine="t", model="t", params={}, segments=[
        seg(0, 0, 2, "A", "First part."), seg(1, 4, 6, "A", "Second part."), seg(2, 7, 8, "B", "Reply.")])


def test_merge_turns_option_joins_consecutive_same_speaker_lines():
    t = _two_turns()
    merged = render("txt", t, None, ExportOptions(merge_turns=True))[0].decode()
    assert "[00:00:00.0] A: First part. Second part." in merged and merged.count("A:") == 1
    split = render("txt", t, None, ExportOptions(merge_turns=False))[0].decode()
    assert split.count("A:") == 2
    # captions and CSV always keep timed cues, whatever the option says
    assert render("vtt", t, None, ExportOptions(merge_turns=True))[0].decode().count("A: ") == 2
    assert render("csv", t, None, ExportOptions(merge_turns=True))[0].decode().count(",A,") == 2


def test_merge_turns_in_word_table():
    t = _two_turns()
    doc = Document(io.BytesIO(render("docx_table", t, None, ExportOptions(merge_turns=True))[0]))
    assert len(doc.tables[0].rows) == 3  # header + merged A + B
    assert doc.tables[0].rows[1].cells[0].text == "00:00:00.0-00:00:06.0"
    default = Document(io.BytesIO(render("docx_table", t, None)[0]))
    assert len(default.tables[0].rows) == 4  # unchanged default: one row per segment
